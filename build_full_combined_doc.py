import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def generate_full_document():
    # Load base document
    base_file = r"g:\BTLWNC\FlawFirst.Api\Phan_Tich_Thiet_Ke_He_Thong_FlawFirst.docx"
    doc = Document(base_file)

    # Styles
    FONT_FAMILY = "Times New Roman"
    COLOR_PRIMARY = RGBColor(0x0F, 0x17, 0x2A)     # Dark Slate
    COLOR_TITLE = RGBColor(0x1E, 0x3A, 0x8A)       # Navy
    COLOR_SUB = RGBColor(0x03, 0x69, 0xA1)         # Ocean
    COLOR_TEXT = RGBColor(0x1E, 0x29, 0x3B)        # Slate 800
    COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)       # Slate 500

    HEX_HEADER_BG = "1E3A8A"                       # Dark Navy
    HEX_ZEBRA = "F8FAFC"                           # Soft Gray
    HEX_BORDER = "CBD5E1"                          # Slate Border

    def format_run(run, font_name=FONT_FAMILY, size_pt=12, bold=False, italic=False, color=COLOR_TEXT):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=2, space_after=6, line_spacing=1.25):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            format_run(r, FONT_FAMILY, 12, False, False, COLOR_TEXT)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, FONT_FAMILY, 15, True, False, COLOR_TITLE)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, FONT_FAMILY, 13.5, True, False, COLOR_TITLE)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, FONT_FAMILY, 12, True, False, COLOR_PRIMARY)
        return p

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(m)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_table_title(title_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        format_run(r, FONT_FAMILY, 11.5, True, False, COLOR_PRIMARY)
        return p

    def add_image_title(title_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(14)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        format_run(r, FONT_FAMILY, 11, False, True, COLOR_PRIMARY)
        return p

    # Add page break before starting Database Design section
    doc.add_page_break()

    # ----------------------------------------------------
    # PHẦN 4: THIẾT KẾ CƠ SỞ DỮ LIỆU & BIỂU ĐỒ ERD
    # ----------------------------------------------------
    add_h1("PHẦN 4: THIẾT KẾ CƠ SỞ DỮ LIỆU VÀ BIỂU ĐỒ QUAN HỆ THỰC THỂ (ERD)")
    add_h2("3.1.2. Phân tích hệ thống về dữ liệu")
    add_h3("a. Xác định thực thể (Entity) và mối quan hệ")

    add_p(
        "Dựa trên các yêu cầu nghiệp vụ và các quy trình khảo sát thực tế của dự án FlawFirst, "
        "hệ thống được thiết kế gồm 11 thực thể (Entities). Các thực thể này được phân chia thành hai nhóm chính: "
        "nhóm thực thể nghiệp vụ cốt lõi (lưu trữ thông tin người dùng, sản phẩm công nghệ, bài viết đánh giá bóc tách ưu và nhược điểm, "
        "các tương tác xã hội như bình luận, bình chọn hữu ích, đánh dấu lưu bài, theo dõi tác giả, cùng phân hệ xử lý báo cáo vi phạm) "
        "và nhóm thực thể danh mục & phụ trợ (quản lý phân loại ngành hàng, lưu trữ hình ảnh minh chứng thực tế của bài đánh giá, "
        "hình ảnh bằng chứng tố cáo vi phạm). Các thực thể nghiệp vụ chính và các thực thể danh mục phụ trợ được mô tả chi tiết trong Bảng 3.10."
    )

    # Bảng 3.10
    add_table_title("Bảng 3.10. Các thực thể của hệ thống")
    tbl_entities = doc.add_table(rows=12, cols=4)
    tbl_entities.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_entities.autofit = False
    set_table_borders(tbl_entities, color=HEX_BORDER, sz="4", val="single")

    ent_headers = ["STT", "Thực thể", "Nhóm", "Mô tả vai trò nghiệp vụ"]
    ent_widths = [Inches(0.6), Inches(1.6), Inches(1.4), Inches(2.9)]

    for c_idx, h_text in enumerate(ent_headers):
        cell = tbl_entities.cell(0, c_idx)
        cell.width = ent_widths[c_idx]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        format_run(r, FONT_FAMILY, 10.5, True, False, RGBColor(0xFF, 0xFF, 0xFF))

    entities_data = [
        ("1", "User\n(Người dùng)", "Nghiệp vụ", "Tài khoản người dùng trong hệ thống: lưu trữ họ tên, email định danh (duy nhất), mật khẩu băm mã hóa BCrypt, vai trò phân quyền (Admin, User), trạng thái khóa tài khoản và thời điểm tạo."),
        ("2", "Category\n(Danh mục)", "Danh mục", "Phân loại ngành hàng của sản phẩm (ví dụ: Điện thoại, Laptop, Máy tính bảng, Tai nghe...): tên danh mục là duy nhất, kèm theo mô tả tóm tắt."),
        ("3", "Product\n(Sản phẩm)", "Nghiệp vụ", "Thông tin sản phẩm công nghệ cần được cộng đồng đánh giá: tên sản phẩm, thương hiệu sản xuất, mô tả cấu hình, mức giá niêm yết, đường dẫn ảnh đại diện và liên kết tới danh mục ngành hàng."),
        ("4", "Review\n(Đánh giá Flaw-First)", "Nghiệp vụ", "Thực thể trọng tâm mang tính đột phá của hệ thống: bắt buộc bóc tách rành mạch giữa Ưu điểm (Pros) và Nhược điểm thực tế (Cons), điểm xếp hạng (1–5 sao), cờ ẩn danh, cờ kiểm duyệt ẩn bài, bộ đếm Upvote/Downvote."),
        ("5", "ReviewImage\n(Ảnh đánh giá)", "Phụ trợ", "Lưu trữ danh sách hình ảnh thực tế do người đánh giá chụp và tải lên kèm theo bài review (tối đa 5 ảnh / bài), phục vụ minh chứng cho các điểm ưu/nhược điểm đã nêu."),
        ("6", "Comment\n(Bình luận)", "Nghiệp vụ", "Lưu trữ các phản hồi, trao đổi, thảo luận của cộng đồng dưới mỗi bài viết đánh giá; hỗ trợ tùy chọn bình luận ẩn danh để bảo vệ quyền riêng tư."),
        ("7", "ReviewVote\n(Bình chọn)", "Nghiệp vụ", "Ghi nhận lượt đánh giá chất lượng bài review: người dùng bình chọn 'Hữu ích' (Upvote) hoặc 'Không hữu ích' (Downvote); mỗi người dùng chỉ được vote duy nhất 1 lần trên 1 review."),
        ("8", "Bookmark\n(Lưu đánh giá)", "Nghiệp vụ", "Lưu trữ danh sách các bài đánh giá tâm đắc mà người dùng đánh dấu để xem lại trong thư viện cá nhân; tránh trùng lặp thông qua ràng buộc duy nhất (ReviewId, UserId)."),
        ("9", "UserFollow\n(Theo dõi)", "Nghiệp vụ", "Mạng lưới tương tác xã hội giữa các thành viên: ghi nhận quan hệ người theo dõi (Follower) và người được theo dõi (Following); ràng buộc không cho phép tự theo dõi chính mình."),
        ("10", "ReviewReport\n(Báo cáo vi phạm)", "Nghiệp vụ", "Quản lý quy trình kiểm duyệt nội dung: lưu trữ tố cáo từ người dùng về bài review vi phạm chính sách (Spam, xúc phạm, sai lệch), ảnh chụp nội dung tại thời điểm tố cáo, trạng thái và quyết định của Admin."),
        ("11", "ReportEvidence\n(Bằng chứng tố cáo)", "Phụ trợ", "Lưu trữ các hình ảnh chụp màn hình bằng chứng vi phạm do người tố cáo đính kèm khi gửi báo cáo vi phạm lên Quản trị viên.")
    ]

    for r_idx, row_vals in enumerate(entities_data):
        row_cells = tbl_entities.rows[r_idx + 1].cells
        bg_color = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row_cells[c_idx]
            cell.width = ent_widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            format_run(r, FONT_FAMILY, 10, (c_idx == 1), False, COLOR_TEXT)

    add_p("", space_after=4)

    # Bảng 3.11: Mối quan hệ
    add_p(
        "Các mối quan hệ giữa các thực thể trong cơ sở dữ liệu FlawFirst được tổng hợp trong Bảng 3.11 "
        "(ký hiệu 1 – N: một – nhiều; 1 – 1: một – một). Mỗi mối quan hệ đều được thiết lập chặt chẽ thông qua các "
        "khóa ngoại (Foreign Key) cùng các quy tắc hành vi toàn vẹn dữ liệu (Cascade, Restrict, SetNull) "
        "nhằm đảm bảo tính nhất quán tuyệt đối trong suốt vòng đời vận hành hệ thống."
    )

    add_table_title("Bảng 3.11. Các mối quan hệ giữa các thực thể")
    tbl_relations = doc.add_table(rows=17, cols=4)
    tbl_relations.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_relations.autofit = False
    set_table_borders(tbl_relations, color=HEX_BORDER, sz="4", val="single")

    rel_headers = ["Thực thể cha", "Bản số", "Thực thể con", "Ý nghĩa nghiệp vụ & Ràng buộc toàn vẹn"]
    rel_widths = [Inches(1.5), Inches(0.9), Inches(1.5), Inches(2.6)]

    for c_idx, h_text in enumerate(rel_headers):
        cell = tbl_relations.cell(0, c_idx)
        cell.width = rel_widths[c_idx]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, 100, 100, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        format_run(r, FONT_FAMILY, 10.5, True, False, RGBColor(0xFF, 0xFF, 0xFF))

    relations_data = [
        ("Category", "1 – N", "Product", "Một danh mục chứa nhiều sản phẩm; mỗi sản phẩm thuộc đúng một danh mục. Ràng buộc Restrict (không được xóa danh mục khi còn sản phẩm trực thuộc)."),
        ("Product", "1 – N", "Review", "Một sản phẩm có nhiều bài đánh giá từ cộng đồng; mỗi bài đánh giá chỉ dành cho một sản phẩm. Ràng buộc Restrict; Unique Index (ProductId, UserId) đảm bảo 1 người chỉ review 1 lần."),
        ("User", "1 – N", "Review", "Một người dùng có thể viết nhiều bài review cho các sản phẩm khác nhau; mỗi bài review do đúng một người dùng tạo. Ràng buộc Restrict."),
        ("Review", "1 – N", "ReviewImage", "Một bài review có thể đính kèm nhiều hình ảnh minh chứng thực tế (0–5 ảnh); mỗi ảnh thuộc về đúng một bài review. Ràng buộc Cascade Delete (xóa review thì ảnh tự động bị xóa)."),
        ("Review", "1 – N", "Comment", "Một bài review có nhiều bình luận thảo luận; mỗi bình luận thuộc về một bài review. Ràng buộc Cascade Delete."),
        ("User", "1 – N", "Comment", "Một người dùng có thể đăng nhiều bình luận; mỗi bình luận do đúng một người dùng thực hiện. Ràng buộc Restrict."),
        ("Review", "1 – N", "ReviewVote", "Một bài review nhận nhiều lượt bình chọn; mỗi lượt bình chọn ghi nhận cho một review. Ràng buộc Cascade Delete; Unique Index (ReviewId, UserId) ngăn vote nhiều lần."),
        ("User", "1 – N", "ReviewVote", "Một người dùng có thể vote cho nhiều bài review khác nhau; mỗi lượt vote do một người dùng thực hiện. Ràng buộc Restrict."),
        ("Review", "1 – N", "Bookmark", "Một bài review có thể được nhiều người dùng đánh dấu lưu lại; mỗi bản ghi bookmark trỏ đến một review. Ràng buộc Cascade Delete."),
        ("User", "1 – N", "Bookmark", "Một người dùng có thể lưu nhiều bài review vào bộ sưu tập cá nhân. Ràng buộc Restrict; Unique Index (ReviewId, UserId) tránh lưu trùng lặp."),
        ("User", "1 – N", "UserFollow (Follower)", "Một người dùng có thể theo dõi nhiều tác giả khác nhau (danh sách Following). Ràng buộc Restrict; CHECK FollowerId <> FollowingId (không được tự theo dõi chính mình)."),
        ("User", "1 – N", "UserFollow (Following)", "Một người dùng có thể được theo dõi bởi nhiều người khác (danh sách Followers). Ràng buộc Restrict; Unique Index (FollowerId, FollowingId)."),
        ("Review", "1 – N", "ReviewReport", "Một bài review có thể nhận nhiều lượt báo cáo vi phạm; mỗi báo cáo gắn liền với một review. Ràng buộc SetNull (khi bài review bị xóa vĩnh viễn, ReviewId đổi thành NULL để lưu lại lịch sử xử lý của Admin)."),
        ("User", "1 – N", "ReviewReport (Reporter)", "Một người dùng có thể gửi nhiều báo cáo vi phạm cho các bài review vi phạm chính sách; mỗi báo cáo ghi nhận người tố cáo. Ràng buộc Restrict."),
        ("User", "1 – N", "ReviewReport (ResolvedAdmin)", "Một Quản trị viên (Admin) có thể thụ lý và xử lý nhiều báo cáo vi phạm; báo cáo đã giải quyết ghi nhận Admin phụ trách. Ràng buộc Restrict."),
        ("ReviewReport", "1 – N", "ReportEvidence", "Một báo cáo vi phạm có thể đính kèm nhiều hình ảnh chụp màn hình làm bằng chứng; mỗi ảnh bằng chứng thuộc về một báo cáo. Ràng buộc Cascade Delete.")
    ]

    for r_idx, row_vals in enumerate(relations_data):
        row_cells = tbl_relations.rows[r_idx + 1].cells
        bg_color = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row_cells[c_idx]
            cell.width = rel_widths[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 2] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            is_bold = (c_idx in [0, 2])
            format_run(r, FONT_FAMILY, 9.5, is_bold, False, COLOR_TEXT)

    doc.add_page_break()

    # Mục b: ERD
    add_h3("b. Biểu đồ quan hệ thực thể — ERD (mô tả quan hệ)")
    add_p(
        "Sơ đồ ERD thể hiện mối liên kết logic giữa các thực thể, đảm bảo tính toàn vẹn dữ liệu, "
        "phản ánh chính xác luồng nghiệp vụ đánh giá Flaw-First và được chuẩn hóa ở mức 3NF (Third Normal Form). "
        "Trong biểu đồ, thực thể Review đóng vai trò là thực thể trung tâm của toàn bộ hệ thống, kết nối trực tiếp "
        "với Sản phẩm (Product), Người đánh giá (User), đồng thời phân nhánh tới toàn bộ các tương tác xoay quanh bài viết "
        "như Hình ảnh minh chứng (ReviewImage), Bình luận (Comment), Bình chọn hữu ích (ReviewVote), Đánh dấu lưu bài (Bookmark) "
        "và Phân hệ kiểm duyệt tố cáo vi phạm (ReviewReport, ReportEvidence)."
    )

    erd_img_path = r"g:\BTLWNC\FlawFirst.Api\erd_flawfirst.png"
    if os.path.exists(erd_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(erd_img_path, width=Inches(6.5))

    add_image_title("Hình 3.2. Biểu đồ quan hệ thực thể (ERD) của hệ thống FlawFirst")

    # Mục c: Chuẩn hóa
    add_h3("c. Đánh giá mức độ chuẩn hóa dữ liệu của hệ thống")
    add_p(
        "Cơ sở dữ liệu của hệ thống FlawFirst được thiết kế tuân thủ nghiêm ngặt các quy tắc lý thuyết cơ sở dữ liệu quan hệ, "
        "đạt mức độ chuẩn hóa Dạng chuẩn 3 (3NF — Third Normal Form), đồng thời áp dụng kỹ thuật Bán chuẩn hóa có kiểm soát "
        "(Controlled Denormalization) tại các điểm nghẽn hiệu năng trọng yếu nhằm tối ưu hóa tốc độ tải trang trong môi trường thực tế."
    )

    add_p(
        "1. Dạng chuẩn 1 (1NF): Toàn bộ thuộc tính đều mang giá trị nguyên tố, không có mảng hay nhóm lặp. "
        "Danh sách hình ảnh đánh giá và hình ảnh bằng chứng được tách thành các thực thể ReviewImages và ReportEvidences riêng biệt. "
        "Mọi bảng đều có khóa chính Id định danh duy nhất.\n\n"
        "2. Dạng chuẩn 2 (2NF): Đạt 1NF và toàn bộ các bảng đều sử dụng khóa chính đơn (Surrogate Key: Id INT IDENTITY). "
        "Mọi thuộc tính không khóa đều phụ thuộc hàm toàn phần vào khóa chính Id. Các khóa tự nhiên phức hợp như (ProductId, UserId) "
        "được thiết lập Unique Index để bảo đảm nghiệp vụ mà không phá vỡ 2NF.\n\n"
        "3. Dạng chuẩn 3 (3NF): Đạt 2NF và không tồn tại bất kỳ phụ thuộc bắc cầu nào giữa các thuộc tính không khóa. "
        "Bảng Product chỉ lưu CategoryId, bảng Review chỉ lưu ProductId và UserId, bảng Comment chỉ lưu ReviewId và UserId.\n\n"
        "4. Kỹ thuật Bán chuẩn hóa có kiểm soát (Controlled Denormalization):\n"
        "   • Lưu trữ UpvoteCount và DownvoteCount trong bảng Reviews: Tránh việc thực hiện COUNT(*) liên tục trên bảng ReviewVotes "
        "khi hiển thị danh sách đánh giá trên trang sản phẩm, giúp tối ưu hóa truy vấn đọc đạt hiệu năng tức thì O(1).\n"
        "   • Lưu trữ ReviewSnapshot trong bảng ReviewReports: Bảo toàn nguyên vẹn trích đoạn nội dung đánh giá tại thời điểm người dùng gửi tố cáo, "
        "phục vụ kiểm tra đối chiếu ngay cả khi tác giả sửa bài hoặc bài viết bị xóa khỏi hệ thống."
    )

    doc.add_page_break()

    # Mục d: Data Dictionary
    add_h3("d. Từ điển dữ liệu của các bảng (Data Dictionary)")
    add_p(
        "Từ điển dữ liệu (Data Dictionary) mô tả chi tiết cấu trúc lưu trữ vật lý của toàn bộ 11 bảng trong cơ sở dữ liệu FlawFirst. "
        "Mỗi bảng cung cấp đầy đủ: STT, Tên trường (Column Name), Kiểu dữ liệu trong SQL Server & C#, Phân loại khóa (PK / FK), "
        "Khả năng chấp nhận giá trị Null, Ràng buộc / Mặc định và Diễn giải ý nghĩa nghiệp vụ cụ thể."
    )

    def render_dict_sub(table_num, table_name, table_desc, columns_data):
        add_table_title(f"Bảng {table_num}. Từ điển dữ liệu bảng {table_name}")
        add_p(f"• Mục đích: {table_desc}", space_before=1, space_after=4)

        tbl = doc.add_table(rows=len(columns_data) + 1, cols=7)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl, color=HEX_BORDER, sz="4", val="single")

        dict_headers = ["STT", "Tên trường", "Kiểu dữ liệu", "Khóa", "Null", "Ràng buộc / Mặc định", "Diễn giải ý nghĩa"]
        dict_widths = [Inches(0.45), Inches(1.2), Inches(1.2), Inches(0.6), Inches(0.55), Inches(1.3), Inches(1.5)]

        for c_idx, h_text in enumerate(dict_headers):
            cell = tbl.cell(0, c_idx)
            cell.width = dict_widths[c_idx]
            set_cell_background(cell, HEX_HEADER_BG)
            set_cell_margins(cell, 80, 80, 90, 90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            format_run(r, FONT_FAMILY, 9.5, True, False, RGBColor(0xFF, 0xFF, 0xFF))

        for r_idx, row_vals in enumerate(columns_data):
            row_cells = tbl.rows[r_idx + 1].cells
            bg_color = HEX_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_vals):
                cell = row_cells[c_idx]
                cell.width = dict_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, 70, 70, 80, 80)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                is_bold = (c_idx == 1 or "PK" in val)
                text_color = RGBColor(0x1E, 0x40, 0xAF) if "PK" in val else (RGBColor(0x9A, 0x34, 0x12) if "FK" in val else COLOR_TEXT)
                format_run(r, FONT_FAMILY, 9, is_bold, False, text_color)

        add_p("", space_after=6)

    # 11 Tables
    dict_tables_all = [
        ("3.12", "Users", "Lưu trữ danh sách tài khoản định danh người dùng trong hệ thống FlawFirst, phân quyền truy cập và kiểm soát trạng thái khóa.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng, định danh duy nhất mỗi người dùng."),
            ("2", "FullName", "nvarchar(150) (string)", "-", "No", "Not Empty", "Họ và tên đầy đủ hiển thị trên hệ thống."),
            ("3", "Email", "nvarchar(256) (string)", "-", "No", "UNIQUE INDEX", "Địa chỉ email duy nhất, dùng làm tên đăng nhập."),
            ("4", "PasswordHash", "nvarchar(200) (string)", "-", "No", "BCrypt Hashed", "Chuỗi mật khẩu đã được băm an toàn bằng thuật toán BCrypt."),
            ("5", "Role", "nvarchar(20) (string)", "-", "No", "CHECK in ('Admin','User')", "Vai trò tài khoản: 'Admin' (Quản trị viên) hoặc 'User' (Thành viên)."),
            ("6", "IsLocked", "bit (bool)", "-", "No", "Default 0 (false)", "Trạng thái khóa tài khoản: 1 = Bị khóa bởi Admin, 0 = Đang hoạt động."),
            ("7", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm tạo tài khoản trên hệ thống.")
        ]),
        ("3.13", "Categories", "Quản lý danh mục ngành hàng sản phẩm, phục vụ phân loại và lọc sản phẩm trên giao diện khám phá.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng, định danh danh mục."),
            ("2", "Name", "nvarchar(120) (string)", "-", "No", "UNIQUE INDEX", "Tên danh mục ngành hàng (ví dụ: Điện thoại, Laptop...)."),
            ("3", "Description", "nvarchar(500) (string)", "-", "Yes", "Null", "Mô tả ngắn gọn về đặc trưng ngành hàng."),
            ("4", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm khởi tạo danh mục.")
        ]),
        ("3.14", "Products", "Lưu trữ thông tin chi tiết về sản phẩm công nghệ cần được cộng đồng đánh giá và bóc tách lỗi.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng, định danh sản phẩm."),
            ("2", "Name", "nvarchar(200) (string)", "-", "No", "Not Empty", "Tên thương mại đầy đủ của sản phẩm."),
            ("3", "Brand", "nvarchar(120) (string)", "-", "No", "Not Empty", "Thương hiệu sản xuất (Apple, Samsung, Dell...)."),
            ("4", "Description", "nvarchar(4000) (string)", "-", "Yes", "Null", "Mô tả chi tiết cấu hình, thông số kỹ thuật sản phẩm."),
            ("5", "Price", "decimal(18,2) (decimal)", "-", "No", "INDEX (Giá tiền)", "Giá bán niêm yết tham khảo tính theo VNĐ."),
            ("6", "ImageUrl", "nvarchar(500) (string)", "-", "Yes", "Null", "Đường dẫn URL ảnh minh họa đại diện của sản phẩm."),
            ("7", "CategoryId", "int (int)", "FK", "No", "FK -> Categories(Id), Restrict", "Mã danh mục ngành hàng mà sản phẩm trực thuộc."),
            ("8", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm thêm sản phẩm vào hệ thống.")
        ]),
        ("3.15", "Reviews", "Thực thể trung tâm lưu trữ bài đánh giá bóc tách ưu và nhược điểm (Flaw-First Review), điểm số và các cờ kiểm duyệt.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng, định danh bài đánh giá."),
            ("2", "ProductId", "int (int)", "FK", "No", "FK -> Products(Id), Restrict", "Mã sản phẩm được đánh giá."),
            ("3", "UserId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng viết bài đánh giá."),
            ("4", "Pros", "nvarchar(4000) (string)", "-", "No", "Not Empty", "Nội dung các điểm hài lòng, ưu điểm thực tế."),
            ("5", "Cons", "nvarchar(4000) (string)", "-", "No", "Not Empty", "Nội dung các nhược điểm, hạn chế cốt lõi (Flaw-First)."),
            ("6", "Rating", "int (int)", "-", "No", "CHECK [Rating] 1..5", "Số sao đánh giá tổng quan (từ 1 đến 5 sao)."),
            ("7", "IsAnonymous", "bit (bool)", "-", "No", "Default 0 (false)", "Tùy chọn hiển thị ẩn danh tên tác giả với cộng đồng."),
            ("8", "IsHidden", "bit (bool)", "-", "No", "Default 0 (false)", "Cờ ẩn bài viết khỏi trang công khai bởi Admin."),
            ("9", "UpvoteCount", "int (int)", "-", "No", "Default 0 (Denormalized)", "Tổng số lượt bình chọn Hữu ích."),
            ("10", "DownvoteCount", "int (int)", "-", "No", "Default 0 (Denormalized)", "Tổng số lượt bình chọn Không hữu ích."),
            ("11", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm đăng bài đánh giá."),
            ("12", "UpdatedAt", "datetime2 (DateTime)", "-", "Yes", "Null", "Thời điểm cập nhật chỉnh sửa bài review lần cuối.")
        ]),
        ("3.16", "ReviewImages", "Lưu trữ đường dẫn các tệp ảnh thực tế đính kèm trong bài review do người dùng tải lên để làm minh chứng.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của bản ghi hình ảnh."),
            ("2", "ReviewId", "int (int)", "FK", "No", "FK -> Reviews(Id), Cascade", "Mã bài review sở hữu hình ảnh minh chứng."),
            ("3", "ImageUrl", "nvarchar(500) (string)", "-", "No", "URL format", "Đường dẫn tệp ảnh lưu trữ trên máy chủ (/uploads/reviews/...).")
        ]),
        ("3.17", "Comments", "Lưu trữ bình luận, trao đổi, hỏi đáp giữa các thành viên cộng đồng dưới mỗi bài viết đánh giá.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của bình luận."),
            ("2", "ReviewId", "int (int)", "FK", "No", "FK -> Reviews(Id), Cascade", "Mã bài review được bình luận."),
            ("3", "UserId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng gửi bình luận."),
            ("4", "Content", "nvarchar(2000) (string)", "-", "No", "Not Empty", "Nội dung văn bản trao đổi của bình luận."),
            ("5", "IsAnonymous", "bit (bool)", "-", "No", "Default 0 (false)", "Tùy chọn ẩn danh tác giả bình luận."),
            ("6", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm gửi bình luận lên hệ thống.")
        ]),
        ("3.18", "ReviewVotes", "Ghi nhận tương tác bình chọn hữu ích (Upvote) hoặc không hữu ích (Downvote) của từng người dùng trên mỗi bài review.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của lượt bình chọn."),
            ("2", "ReviewId", "int (int)", "FK", "No", "FK -> Reviews(Id), Cascade", "Mã bài review nhận lượt bình chọn."),
            ("3", "UserId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng thực hiện bình chọn."),
            ("4", "IsUpvote", "bit (bool)", "-", "No", "1: Upvote, 0: Downvote", "Loại bình chọn: True = Hữu ích, False = Không hữu ích."),
            ("5", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm thực hiện hành động vote.")
        ]),
        ("3.19", "Bookmarks", "Lưu trữ danh sách bài viết đánh giá mà người dùng đánh dấu để lưu trữ và xem lại trong bộ sưu tập cá nhân.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của bản ghi lưu bài."),
            ("2", "ReviewId", "int (int)", "FK", "No", "FK -> Reviews(Id), Cascade", "Mã bài review được đánh dấu lưu."),
            ("3", "UserId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng thực hiện lưu bài."),
            ("4", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm thực hiện đánh dấu lưu bài.")
        ]),
        ("3.20", "UserFollows", "Mô hình hóa quan hệ theo dõi tác giả (Follower / Following) giữa các thành viên trên mạng xã hội FlawFirst.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của quan hệ theo dõi."),
            ("2", "FollowerId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng nhấn theo dõi (Người theo dõi)."),
            ("3", "FollowingId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng được theo dõi (Tác giả được follow)."),
            ("4", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm bắt đầu mối quan hệ theo dõi.")
        ]),
        ("3.21", "ReviewReports", "Lưu trữ các báo cáo vi phạm nội dung từ cộng đồng, snapshot bài viết và quyết định xử lý kiểm duyệt của Quản trị viên.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của báo cáo vi phạm."),
            ("2", "ReviewId", "int (int)", "FK", "Yes", "FK -> Reviews(Id), SetNull", "Mã bài review bị báo cáo (chuyển NULL nếu review bị xóa)."),
            ("3", "ReporterId", "int (int)", "FK", "No", "FK -> Users(Id), Restrict", "Mã người dùng gửi báo cáo vi phạm."),
            ("4", "Reason", "nvarchar(2000) (string)", "-", "No", "Not Empty", "Lý do và mô tả chi tiết hành vi vi phạm."),
            ("5", "ReviewSnapshot", "nvarchar(1000) (string)", "-", "Yes", "Null", "Trích đoạn nội dung review tại thời điểm báo cáo vi phạm."),
            ("6", "Status", "nvarchar(20) (string)", "-", "No", "CHECK in ('Pending','Resolved','Dismissed')", "Trạng thái báo cáo: Chờ xử lý, Đã giải quyết, hoặc Bị bác bỏ."),
            ("7", "Action", "nvarchar(20) (string)", "-", "Yes", "Null", "Hành động xử lý của Admin ('Hide', 'Delete', 'Ban', 'Dismiss')."),
            ("8", "ResolvedByAdminId", "int (int)", "FK", "Yes", "FK -> Users(Id), Restrict", "Mã Quản trị viên đã thụ lý và giải quyết báo cáo."),
            ("9", "CreatedAt", "datetime2 (DateTime)", "-", "No", "Default UTC Now", "Thời điểm gửi báo cáo vi phạm."),
            ("10", "ResolvedAt", "datetime2 (DateTime)", "-", "Yes", "Null", "Thời điểm Admin hoàn tất xử lý báo cáo.")
        ]),
        ("3.22", "ReportEvidences", "Lưu trữ các hình ảnh chụp màn hình làm bằng chứng vi phạm do người tố cáo tải lên đính kèm báo cáo.", [
            ("1", "Id", "int (int)", "PK", "No", "IDENTITY(1,1)", "Khóa chính tự tăng của bằng chứng ảnh."),
            ("2", "ReportId", "int (int)", "FK", "No", "FK -> ReviewReports(Id), Cascade", "Mã báo cáo vi phạm mà hình ảnh thuộc về."),
            ("3", "ImageUrl", "nvarchar(500) (string)", "-", "No", "URL format", "Đường dẫn tệp ảnh bằng chứng trên máy chủ (/uploads/evidences/...).")
        ])
    ]

    for t_num, t_name, t_desc, cols in dict_tables_all:
        render_dict_sub(t_num, t_name, t_desc, cols)

    out_file = r"g:\BTLWNC\FlawFirst.Api\Phan_Tich_Thiet_Ke_He_Thong_FlawFirst_Full.docx"
    doc.save(out_file)
    print("Full combined document generated successfully at:", out_file)

if __name__ == "__main__":
    generate_full_document()
